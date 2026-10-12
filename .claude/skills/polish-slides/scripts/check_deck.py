#!/usr/bin/env python3
"""Render a .pptx and report layout problems a human would spot.

Usage: python check_deck.py deck.pptx [outdir]

Writes contact sheets (sheet-N.jpg, 6 slides each) and per-slide PNGs to outdir,
then prints findings: text outside the slide, tiny fonts, leftover placeholder
text, slides with almost no content, and words broken across two lines (a
narrow box or chevron splitting "Regression" into "Regressio" / "n"). Findings
are hints, not verdicts: always look at the contact sheets too.

Run it with a Python that has pymupdf and Pillow; the broken-word check also
needs python-pptx (skipped with a note if it is missing).
"""
import re, subprocess, sys, shutil, pathlib
import pymupdf
from PIL import Image

PLACEHOLDER = re.compile(r"lorem|ipsum|sample (text|footer)|click to (edit|add)|\bTODO\b|\bxxx+\b|\[insert|Typoblind|Muster", re.I)
FOOTER_ZONE = 0.92      # text below this fraction of slide height is footer/page number
MIN_BODY_PT = 10.5      # below this, text is hard to read when projected
MIN_CODE_PT = 10.0      # monospaced code blocks may go down to 10 pt
TOLERANCE = 0.05        # the PDF reports 10.5 pt as 10.49
MONO = re.compile(r"mono|courier|consol", re.I)
CARD = re.compile(r"Welthandelsplatz")   # template business card on the closing slide: fixed small text
WORD = re.compile(r"[A-Za-zÀ-ÿ]+")


def deck_words(deck):
    """All words typed in the deck, to tell a word broken by wrapping from two real words."""
    try:
        from pptx import Presentation
    except ImportError:
        return None
    words = set()
    for slide in Presentation(deck).slides:
        for sh in slide.shapes:
            frames = [sh.text_frame] if sh.has_text_frame else []
            if sh.has_table:
                frames += [c.text_frame for r in sh.table.rows for c in r.cells]
            for tf in frames:
                words.update(w.lower() for w in WORD.findall(tf.text))
    return words


def broken_words(page, words):
    """Consecutive lines whose last and first fragments only form a word together.

    LibreOffice may put the two halves in separate text blocks, so the lines of the
    whole page are compared in reading order."""
    lines = [ln for b in page.get_text("dict")["blocks"] for l in b.get("lines", [])
             if (ln := "".join(s["text"] for s in l["spans"]).strip())]
    found = []
    for a, c in zip(lines, lines[1:]):
        ta, tc = WORD.findall(a), WORD.findall(c)
        if not ta or not tc or not a[-1].isalpha() or not c[0].isalpha():
            continue
        left, right = ta[-1].lower(), tc[0].lower()
        if (left + right) in words and (left not in words or right not in words):
            found.append(ta[-1] + "|" + tc[0])
    return found

def render(deck, out):
    out.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice (soffice) not found: install libreoffice-impress")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(deck)],
                   check=True, capture_output=True, timeout=300)
    pdf = out / (deck.stem + ".pdf")
    if not pdf.exists():
        sys.exit("PDF conversion failed (is the Impress component installed?)")
    return pdf

def main():
    deck = pathlib.Path(sys.argv[1]).resolve()
    out = pathlib.Path(sys.argv[2] if len(sys.argv) > 2 else deck.with_suffix("").name + "_qa").resolve()
    pdf = render(deck, out)
    doc = pymupdf.open(pdf)
    findings, pngs = [], []
    vocab = deck_words(deck)
    if vocab is None:
        findings.append("note: python-pptx not installed, broken-word check skipped")
    for i, page in enumerate(doc, 1):
        W, H = page.rect.width, page.rect.height
        png = out / f"slide-{i:02d}.png"
        page.get_pixmap(dpi=80).save(png); pngs.append(png)
        text = page.get_text()
        if PLACEHOLDER.search(text):
            findings.append(f"slide {i}: leftover placeholder text ({PLACEHOLDER.search(text).group(0)!r})")
        words = len(text.split())
        if words < 6:
            findings.append(f"slide {i}: almost empty ({words} words); fine for a divider, odd otherwise")
        if vocab:
            for frag in broken_words(page, vocab):
                findings.append(f"slide {i}: word broken across lines ({frag}); widen the box or shorten the text")
        d = page.get_text("dict")
        small = set()
        card = bool(CARD.search(text))
        for b in d["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    if not s["text"].strip():
                        continue
                    x0, y0, x1, y1 = s["bbox"]
                    if x1 > W + 1 or y1 > H + 1 or x0 < -1 or y0 < -1:
                        findings.append(f"slide {i}: text runs off the slide: {s['text'][:40]!r}")
                    # pdf points scale with slide width: normalise to a 10 in wide slide
                    pt = s["size"] * (720 / W)
                    limit = MIN_CODE_PT if MONO.search(s["font"]) else MIN_BODY_PT
                    if pt < limit - TOLERANCE and y0 < FOOTER_ZONE * H and not card:
                        small.add(round(pt, 1))
        if small:
            findings.append(f"slide {i}: text smaller than {MIN_BODY_PT} pt (sizes {sorted(small)}); consider splitting the slide")
    for k in range(0, len(pngs), 6):
        ims = [Image.open(p) for p in pngs[k:k + 6]]
        w, h = ims[0].size
        sheet = Image.new("RGB", (w * 2, h * 3), "white")
        for j, im in enumerate(ims):
            sheet.paste(im, ((j % 2) * w, (j // 2) * h))
        sheet.save(out / f"sheet-{k // 6 + 1}.jpg", quality=85)
    print(f"{len(doc)} slides rendered to {out}")
    print("\n".join(findings) if findings else "no automatic findings (still look at the sheets)")

if __name__ == "__main__":
    main()
